"""Word document generator for the Migration Blueprint.

Renders the Pass 3 Blueprint JSON into a corporate Word document focused on
implementation detail — configuration steps, Business Components, TypeScript
extensions, Business Event subscriptions, REST integrations, data migration,
testing, go-live checklist.

Companion to ``doc_generator.py`` (which renders the system documentation /
replaceability analysis). Both run in the QAD-Zone documentation pipeline:
the system doc explains "what does standard QAD already cover", the blueprint
explains "how do you actually build the migration".

Section schema (each section conditional on real data — empty sections skipped):

    Title page → Executive Summary → Migration Strategy + Phases →
    Per-Capability Migration (one section per item, with conditional
    Configuration Steps / Business Component / TypeScript Extension /
    Business Event Subscription / API Integration / Data Migration Notes /
    Risks blocks) → Data Migration (overall) → Testing Plan → Go-Live
    Checklist → References / Citations
"""
from __future__ import annotations

import uuid
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOWNLOADS_DIR = Path("app/static/downloads")

# ── Colour palette — slightly developer-leaning compared to system-doc ───────

C = {
    "BLUE":     RGBColor(0x1F, 0x4E, 0x79),
    "MBLUE":    RGBColor(0x2E, 0x75, 0xB6),
    "INDIGO":   RGBColor(0x44, 0x4C, 0xC4),
    "ACCENT":   RGBColor(0xC0, 0x00, 0x00),
    "MGREEN":   RGBColor(0x70, 0xAD, 0x47),
    "DKGREEN":  RGBColor(0x37, 0x56, 0x23),
    "DGRAY":    RGBColor(0x40, 0x40, 0x40),
    "MGRAY":    RGBColor(0x76, 0x76, 0x76),
    "WHITE":    RGBColor(0xFF, 0xFF, 0xFF),
    "AMBER":    RGBColor(0xC5, 0x5A, 0x11),
    "TEAL":     RGBColor(0x0C, 0x54, 0x60),
    "PURPLE":   RGBColor(0x70, 0x30, 0xA0),
}
FILL = {
    "BLUE":     "1F4E79",
    "MBLUE":    "2E75B6",
    "LBLUE":    "DEEAF1",
    "INDIGO":   "444CC4",
    "LGREEN":   "E2EFDA",
    "LGRAY":    "F2F2F2",
    "WHITE":    "FFFFFF",
    "AMBER":    "FFF3CD",
    "TEAL":     "D1ECF1",
    "CODEBG":   "1E1E1E",   # near-black, for TS code blocks
    "CODELT":   "F5F5F5",
    "ACCENT":   "C00000",
}
FONT      = "Arial"
CODE_FONT = "Cascadia Code"
CODE_FONT_FALLBACK = "Consolas"
ALT  = [FILL["WHITE"], FILL["LGRAY"]]

APPROACH_FILL = {
    "Configuration":          (FILL["LGREEN"], C["DKGREEN"]),
    "TypeScript Extension":   (FILL["INDIGO"], C["WHITE"]),
    "Business Event":         (FILL["LBLUE"],  C["BLUE"]),
    "Hybrid":                 (FILL["AMBER"],  C["AMBER"]),
}


# ── Helpers ───────────────────────────────────────────────────────────────────

def _has(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (list, dict)):
        return len(value) > 0
    if isinstance(value, str):
        v = value.strip()
        return bool(v) and v not in ("", "AUTO")
    return bool(value)


def _shading(cell, fill_hex: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def _borders(cell, color: str = "AAAAAA", size: str = "4") -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcB = OxmlElement("w:tcBorders")
    for side in ["top", "left", "bottom", "right"]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)
        tcB.append(el)
    tcPr.append(tcB)


def _cell_margins(cell) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcM = OxmlElement("w:tcMar")
    for side, val in [("top", 80), ("bottom", 80), ("left", 120), ("right", 120)]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        tcM.append(el)
    tcPr.append(tcM)


def _set_cell(cell, text: str, fill: str = FILL["WHITE"], bold: bool = False,
              italic: bool = False, color: RGBColor | None = None,
              hdr_color: str | None = None, font_size: int = 10) -> None:
    _shading(cell, fill)
    _borders(cell, color=hdr_color or "AAAAAA")
    _cell_margins(cell)
    p = cell.paragraphs[0]
    p.clear()
    run = p.add_run(str(text) if text is not None else "")
    run.font.name = FONT
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color or C["DGRAY"]


def _hdr_cell(cell, text: str, fill: str = FILL["INDIGO"]) -> None:
    _shading(cell, fill)
    _borders(cell, color=fill)
    _cell_margins(cell)
    p = cell.paragraphs[0]
    p.clear()
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = C["WHITE"]


def _add_header_footer(doc: Document, system_name: str, system_full: str) -> None:
    for section in doc.sections:
        hdr = section.header
        hdr.is_linked_to_previous = False
        hp = hdr.paragraphs[0] if hdr.paragraphs else hdr.add_paragraph()
        hp.clear()
        hp.paragraph_format.space_after = Pt(3)
        pPr = hp._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "6")
        bot.set(qn("w:color"), FILL["INDIGO"]); bot.set(qn("w:space"), "1")
        pBdr.append(bot); pPr.append(pBdr)
        r1 = hp.add_run(f"{system_name} — Migration Blueprint")
        r1.font.name = FONT; r1.font.size = Pt(9); r1.font.color.rgb = C["INDIGO"]
        r2 = hp.add_run(f"    {system_full}")
        r2.font.name = FONT; r2.font.size = Pt(9); r2.font.color.rgb = C["MGRAY"]

        ftr = section.footer
        ftr.is_linked_to_previous = False
        fp = ftr.paragraphs[0] if ftr.paragraphs else ftr.add_paragraph()
        fp.clear()
        fp.paragraph_format.space_before = Pt(3)
        pPr2 = fp._p.get_or_add_pPr()
        pBdr2 = OxmlElement("w:pBdr"); top = OxmlElement("w:top")
        top.set(qn("w:val"), "single"); top.set(qn("w:sz"), "4")
        top.set(qn("w:color"), FILL["INDIGO"]); top.set(qn("w:space"), "1")
        pBdr2.append(top); pPr2.append(pBdr2)
        r3 = fp.add_run("CONFIDENTIAL — Migration Blueprint    |    Page ")
        r3.font.name = FONT; r3.font.size = Pt(8); r3.font.color.rgb = C["MGRAY"]
        fldChar1 = OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText"); instr.text = "PAGE"
        fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "end")
        page_r = OxmlElement("w:r")
        page_r.append(fldChar1); page_r.append(instr); page_r.append(fldChar2)
        fp._p.append(page_r)


# ── Element builders ──────────────────────────────────────────────────────────

def _h1(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=1)
    for run in p.runs:
        run.font.name = FONT; run.font.size = Pt(18); run.font.bold = True
        run.font.color.rgb = C["INDIGO"]
    p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(8)


def _h2(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=2)
    for run in p.runs:
        run.font.name = FONT; run.font.size = Pt(14); run.font.bold = True
        run.font.color.rgb = C["MBLUE"]
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "3"); bot.set(qn("w:color"), FILL["MBLUE"])
    pBdr.append(bot); pPr.append(pBdr)
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)


def _h3(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=3)
    for run in p.runs:
        run.font.name = FONT; run.font.size = Pt(12); run.font.bold = True
        run.font.color.rgb = C["DGRAY"]
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)


def _h4(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=4)
    for run in p.runs:
        run.font.name = FONT; run.font.size = Pt(11); run.font.bold = True
        run.font.color.rgb = C["INDIGO"]
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)


def _para(doc: Document, text: str, italic: bool = False,
          color: RGBColor | None = None, size: int = 11) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = FONT; run.font.size = Pt(size)
    run.font.italic = italic; run.font.color.rgb = color or C["DGRAY"]


def _bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text)
    run.font.name = FONT; run.font.size = Pt(11); run.font.color.rgb = C["DGRAY"]


def _num_item(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text)
    run.font.name = FONT; run.font.size = Pt(11); run.font.color.rgb = C["DGRAY"]


def _code_block(doc: Document, code: str, *, language: str = "typescript") -> None:
    """Render a fenced code block with a dark background and monospace font."""
    if not code or not code.strip():
        return

    # Caption
    cap = doc.add_paragraph()
    cap.paragraph_format.space_before = Pt(4); cap.paragraph_format.space_after = Pt(0)
    rcap = cap.add_run(f"  {language}")
    rcap.font.name = FONT; rcap.font.size = Pt(8)
    rcap.font.bold = True; rcap.font.color.rgb = C["MGRAY"]

    # Single-cell table with shaded fill — survives MS Word render better than
    # a plain shaded paragraph for multi-line content.
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    _shading(cell, FILL["CODEBG"])
    _borders(cell, color="333333")
    _cell_margins(cell)

    # Replace cell paragraph with code lines
    cell.paragraphs[0].clear()
    first = True
    for line in code.split("\n"):
        if first:
            p = cell.paragraphs[0]
            first = False
        else:
            p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(0)
        run = p.add_run(line if line else " ")
        run.font.name = CODE_FONT
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xD4, 0xD4, 0xD4)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _note(doc: Document, text: str, fill_key: str = "AMBER",
          color_key: str = "AMBER", icon: str = "⚠ NOTE: ") -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), FILL[fill_key])
    pPr.append(shd)
    r1 = p.add_run(icon); r1.font.name = FONT; r1.font.size = Pt(10)
    r1.font.bold = True; r1.font.color.rgb = C[color_key]
    r2 = p.add_run(text); r2.font.name = FONT; r2.font.size = Pt(10)
    r2.font.color.rgb = C[color_key]


def _data_table(doc: Document, headers: list[str], rows: list[list]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=0, cols=len(headers))
    table.style = "Table Grid"
    hrow = table.add_row()
    for i, h in enumerate(headers):
        _hdr_cell(hrow.cells[i], h)
    for ri, row_data in enumerate(rows):
        row = table.add_row()
        for ci in range(len(headers)):
            val = row_data[ci] if ci < len(row_data) else ""
            _set_cell(row.cells[ci], str(val) if val is not None else "", fill=ALT[ri % 2])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def _kv_table(doc: Document, pairs: list[tuple[str, str]],
              header_fill: str = FILL["INDIGO"]) -> None:
    if not pairs:
        return
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in pairs:
        row = table.add_row()
        _hdr_cell(row.cells[0], label, fill=header_fill)
        _set_cell(row.cells[1], value, fill=FILL["LGRAY"])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def _approach_badge(doc: Document, approach: str) -> None:
    fill, color = APPROACH_FILL.get(approach, (FILL["LGRAY"], C["DGRAY"]))
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    pPr.append(shd)
    run = p.add_run(f"  ► Approach: {approach}")
    run.font.name = FONT; run.font.size = Pt(11)
    run.font.bold = True; run.font.color.rgb = color


def _divider(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "4"); bot.set(qn("w:color"), "CCCCCC")
    pBdr.append(bot); pPr.append(pBdr)


# ── Section builders ──────────────────────────────────────────────────────────

def _build_title_page(doc: Document, T: dict, doc_date: str) -> None:
    sp = doc.add_paragraph(); sp.paragraph_format.space_before = Pt(48)

    if _has(T.get("SYSTEM_NAME")):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(T["SYSTEM_NAME"])
        r.font.name = FONT; r.font.size = Pt(48); r.font.bold = True
        r.font.color.rgb = C["INDIGO"]

    if _has(T.get("SYSTEM_FULL_NAME")):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(T["SYSTEM_FULL_NAME"])
        r.font.name = FONT; r.font.size = Pt(20); r.font.color.rgb = C["MBLUE"]

    div = doc.add_paragraph()
    div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    div.paragraph_format.space_before = Pt(8); div.paragraph_format.space_after = Pt(8)
    pPr = div._p.get_or_add_pPr(); pBdr = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "6"); bot.set(qn("w:color"), FILL["INDIGO"])
    pBdr.append(bot); pPr.append(pBdr)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run("Migration Blueprint")
    r2.font.name = FONT; r2.font.size = Pt(22); r2.font.bold = True
    r2.font.color.rgb = C["INDIGO"]

    sub2 = doc.add_paragraph()
    sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = sub2.add_run("How to migrate to standard QAD Adaptive 2025")
    r3.font.name = FONT; r3.font.size = Pt(13); r3.font.italic = True
    r3.font.color.rgb = C["DGRAY"]

    if _has(T.get("TARGET_PLATFORM")):
        plat = doc.add_paragraph()
        plat.alignment = WD_ALIGN_PARAGRAPH.CENTER
        plat.paragraph_format.space_before = Pt(8)
        r4 = plat.add_run(f"Target: {T['TARGET_PLATFORM']}")
        r4.font.name = FONT; r4.font.size = Pt(11); r4.font.color.rgb = C["MGRAY"]

    doc.add_paragraph().paragraph_format.space_before = Pt(20)

    meta_pairs: list[tuple[str, str]] = []
    for label, key in [
        ("System",          "SYSTEM_FULL_NAME"),
        ("Target Platform", "TARGET_PLATFORM"),
        ("Document Type",   "DOCUMENT_TYPE"),
    ]:
        v = T.get(key)
        if _has(v):
            meta_pairs.append((label, v))
    meta_pairs.append(("Document Date", doc_date))
    _kv_table(doc, meta_pairs)

    doc.add_page_break()


def _build_executive_summary(doc: Document, ES: dict) -> None:
    if not _has(ES):
        return
    _h1(doc, "1.  Executive Summary")

    if _has(ES.get("OVERVIEW")):
        _h3(doc, "Migration Overview")
        _para(doc, ES["OVERVIEW"])

    if _has(ES.get("BUSINESS_VALUE")):
        _h3(doc, "Business Value")
        _para(doc, ES["BUSINESS_VALUE"])

    fast_facts: list[tuple[str, str]] = []
    if _has(ES.get("ESTIMATED_EFFORT")):
        fast_facts.append(("Estimated Effort", ES["ESTIMATED_EFFORT"]))
    if _has(ES.get("KEY_DEPENDENCIES")):
        deps = ES["KEY_DEPENDENCIES"]
        fast_facts.append(("Key Dependencies",
                           ", ".join(deps) if isinstance(deps, list) else str(deps)))
    if fast_facts:
        _h3(doc, "Fast Facts")
        _kv_table(doc, fast_facts)

    doc.add_page_break()


def _build_strategy(doc: Document, MS: dict) -> None:
    if not _has(MS):
        return
    _h1(doc, "2.  Migration Strategy")

    if _has(MS.get("INTRO_PARA")):
        _para(doc, MS["INTRO_PARA"])

    phases = MS.get("PHASES") or []
    if phases:
        _h2(doc, "2.1  Phased Migration Plan")
        rows = []
        for ph in phases:
            if not isinstance(ph, dict):
                continue
            num = ph.get("PHASE_NUMBER", "")
            name = ph.get("PHASE_NAME", "")
            duration = ph.get("DURATION", "")
            outcome = ph.get("OUTCOME", "")
            activities = ph.get("ACTIVITIES") or []
            act_str = "; ".join(activities) if isinstance(activities, list) else str(activities)
            rows.append([num, name, duration, act_str, outcome])
        if rows:
            _data_table(doc, ["#", "Phase", "Duration", "Activities", "Outcome"], rows)

    doc.add_page_break()


def _build_capability_block(doc: Document, idx: int, cap: dict) -> None:
    """One CAPABILITY_MIGRATIONS entry — many conditional sub-sections."""
    if not isinstance(cap, dict):
        return

    name = cap.get("CAPABILITY", f"Capability {idx}")
    _h2(doc, f"3.{idx}  {name}")

    # Approach badge
    approach = (cap.get("TARGET_APPROACH") or "").strip()
    if _has(approach):
        _approach_badge(doc, approach)

    # Quick KV: current behaviour, target module, effort, deps
    overview: list[tuple[str, str]] = []
    if _has(cap.get("CURRENT_BEHAVIOUR")):
        overview.append(("Current Custom Behaviour", cap["CURRENT_BEHAVIOUR"]))
    if _has(cap.get("QAD_MODULE")):
        overview.append(("Target QAD Module / Feature", cap["QAD_MODULE"]))
    if _has(cap.get("EFFORT_ESTIMATE")):
        overview.append(("Effort Estimate", cap["EFFORT_ESTIMATE"]))
    if _has(cap.get("DEPENDENCIES")):
        deps = cap["DEPENDENCIES"]
        overview.append(("Dependencies",
                         ", ".join(deps) if isinstance(deps, list) else str(deps)))
    if overview:
        _kv_table(doc, overview)

    if _has(cap.get("MIGRATION_DETAIL")):
        _h3(doc, "How the Migration Works")
        _para(doc, cap["MIGRATION_DETAIL"])

    # Configuration steps
    cfg_steps = cap.get("CONFIGURATION_STEPS") or []
    if cfg_steps:
        _h3(doc, "Configuration Steps")
        for step in cfg_steps:
            if not isinstance(step, dict):
                continue
            num = step.get("STEP_NUMBER", "")
            title = step.get("TITLE", "")
            desc = step.get("DESCRIPTION", "")
            if _has(title):
                _h4(doc, f"Step {num} — {title}" if _has(num) else title)
            if _has(desc):
                _para(doc, desc)
            for action in (step.get("ACTIONS") or []):
                if _has(action):
                    _num_item(doc, action)
            if _has(step.get("REFERENCE")):
                _para(doc, f"Reference: {step['REFERENCE']}",
                      italic=True, color=C["MGRAY"], size=9)

    # Business Component
    bc = cap.get("BUSINESS_COMPONENT") or {}
    if bc.get("SHOW") and _has(bc.get("NAME")):
        _h3(doc, "Business Component Definition")
        bc_kv: list[tuple[str, str]] = []
        bc_kv.append(("Component Name", bc.get("NAME", "")))
        if _has(bc.get("PURPOSE")):
            bc_kv.append(("Purpose", bc["PURPOSE"]))
        if bc_kv:
            _kv_table(doc, bc_kv)

        props = bc.get("PROPERTIES") or []
        if props:
            rows = []
            for pr in props:
                if not isinstance(pr, dict):
                    continue
                rows.append([
                    pr.get("NAME", ""),
                    pr.get("TYPE", ""),
                    "Yes" if pr.get("REQUIRED") else "No",
                    pr.get("DESCRIPTION", ""),
                ])
            if rows:
                _h4(doc, "Properties")
                _data_table(doc, ["Property", "Type", "Required", "Description"], rows)

        validations = bc.get("VALIDATIONS") or []
        if validations:
            _h4(doc, "Validation Rules")
            for v in validations:
                if _has(v):
                    _bullet(doc, v)

        if _has(bc.get("NOTES")):
            _note(doc, bc["NOTES"])

    # TypeScript Extension
    ts = cap.get("TYPESCRIPT_EXTENSION") or {}
    if ts.get("SHOW") and (ts.get("FILES") or _has(ts.get("PURPOSE"))):
        _h3(doc, "TypeScript Extension Code")
        if _has(ts.get("PURPOSE")):
            _para(doc, ts["PURPOSE"])
        for file_entry in (ts.get("FILES") or []):
            if not isinstance(file_entry, dict):
                continue
            fname = file_entry.get("FILENAME", "")
            purpose = file_entry.get("PURPOSE", "")
            code = file_entry.get("CODE", "")
            notes = file_entry.get("NOTES") or []
            if _has(fname):
                _h4(doc, f"📄 {fname}")
            if _has(purpose):
                _para(doc, purpose, italic=True, color=C["MGRAY"], size=10)
            if _has(code):
                _code_block(doc, code, language="typescript")
            for n in notes:
                if _has(n):
                    _bullet(doc, n)

    # Business Event Subscription
    be = cap.get("BUSINESS_EVENT_SUBSCRIPTION") or {}
    if be.get("SHOW") and _has(be.get("EVENT_NAME")):
        _h3(doc, "Business Event Subscription")
        be_kv = [
            ("Event Name", be.get("EVENT_NAME", "")),
            ("Trigger", be.get("TRIGGER", "")),
            ("Handler Logic", be.get("HANDLER_LOGIC", "")),
        ]
        be_kv = [(k, v) for k, v in be_kv if _has(v)]
        if be_kv:
            _kv_table(doc, be_kv)
        payload = be.get("PAYLOAD_FIELDS") or []
        if payload:
            _h4(doc, "Payload Fields")
            for f in payload:
                if _has(f):
                    _bullet(doc, f)

    # API Integration
    api = cap.get("API_INTEGRATION") or {}
    if api.get("SHOW") and _has(api.get("ENDPOINT")):
        _h3(doc, "REST API Integration")
        api_kv = [
            ("Endpoint",   api.get("ENDPOINT", "")),
            ("Purpose",    api.get("PURPOSE", "")),
            ("Auth",       api.get("AUTH", "")),
        ]
        api_kv = [(k, v) for k, v in api_kv if _has(v)]
        if api_kv:
            _kv_table(doc, api_kv)
        if _has(api.get("EXAMPLE_PAYLOAD")):
            _h4(doc, "Example Payload")
            _code_block(doc, api["EXAMPLE_PAYLOAD"], language="json")

    # Data Migration Notes (per-capability)
    dm = cap.get("DATA_MIGRATION_NOTES") or {}
    if dm.get("SHOW"):
        _h3(doc, "Data Migration Notes")
        if _has(dm.get("TABLES_AFFECTED")):
            tables = dm["TABLES_AFFECTED"]
            _para(doc,
                  "Tables: " + (", ".join(tables) if isinstance(tables, list) else str(tables)))
        if _has(dm.get("MIGRATION_APPROACH")):
            _para(doc, dm["MIGRATION_APPROACH"])

    # Risks
    risks = cap.get("RISKS") or []
    if risks:
        _h3(doc, "Risks & Considerations")
        for r in risks:
            if isinstance(r, dict):
                txt = r.get("RISK") or r.get("description") or ""
                mit = r.get("MITIGATION", "")
                if _has(txt):
                    _bullet(doc, txt + (f" — Mitigation: {mit}" if _has(mit) else ""))
            elif _has(r):
                _bullet(doc, r)

    _divider(doc)


def _build_capabilities(doc: Document, capabilities: list) -> None:
    if not capabilities:
        return
    _h1(doc, "3.  Per-Capability Migration Plan")
    _para(doc,
          "Each section below documents one capability of the custom system and "
          "the concrete steps, configuration, and code required to deliver the same "
          "business outcome on standard QAD Adaptive 2025.")

    for i, cap in enumerate(capabilities, 1):
        _build_capability_block(doc, i, cap)

    doc.add_page_break()


def _build_data_migration(doc: Document, DM: dict) -> None:
    if not DM or not DM.get("SHOW"):
        return
    _h1(doc, "4.  Data Migration")
    if _has(DM.get("INTRO_PARA")):
        _para(doc, DM["INTRO_PARA"])
    tt = DM.get("TABLES_TABLE") or {}
    if _has(tt.get("rows")) and _has(tt.get("headers")):
        _data_table(doc, tt["headers"], tt["rows"])
    doc.add_page_break()


def _build_testing(doc: Document, TP: dict) -> None:
    if not _has(TP):
        return
    _h1(doc, "5.  Testing Plan")
    if _has(TP.get("INTRO_PARA")):
        _para(doc, TP["INTRO_PARA"])

    test_cases = TP.get("TEST_CASES") or []
    rows = []
    for i, tc in enumerate(test_cases, 1):
        if not isinstance(tc, dict):
            continue
        scen = tc.get("SCENARIO", "")
        steps = tc.get("STEPS") or []
        steps_str = "; ".join(steps) if isinstance(steps, list) else str(steps)
        expected = tc.get("EXPECTED", "")
        rows.append([str(i), scen, steps_str, expected])
    if rows:
        _data_table(doc, ["#", "Scenario", "Steps", "Expected Result"], rows)

    doc.add_page_break()


def _build_golive(doc: Document, items: list) -> None:
    if not items:
        return
    _h1(doc, "6.  Go-Live Checklist")
    for item in items:
        if _has(item):
            _bullet(doc, item)
    doc.add_page_break()


def _build_references(doc: Document, refs: list) -> None:
    if not refs:
        return
    _h1(doc, "7.  References & Citations")
    rows = []
    for r in refs:
        if not isinstance(r, dict):
            continue
        rows.append([
            r.get("source_doc", ""),
            r.get("section", "") or r.get("label", ""),
        ])
    if rows:
        _data_table(doc, ["Source", "Section"], rows)


def _build_end_page(doc: Document, system_name: str, doc_date: str) -> None:
    _divider(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("— End of Migration Blueprint —")
    r.font.name = FONT; r.font.size = Pt(10); r.font.italic = True; r.font.color.rgb = C["MGRAY"]
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(f"{system_name}  |  Migration Blueprint  |  {doc_date}")
    r2.font.name = FONT; r2.font.size = Pt(9); r2.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)


# ── Main entry point ──────────────────────────────────────────────────────────

def generate_blueprint_document(blueprint: dict, *, system_full: str = "") -> str:
    """Render the Pass 3 blueprint JSON into a Word document.

    Returns the URL path under /static/downloads/ for the saved .docx.
    """
    DOWNLOADS_DIR.mkdir(parents=True, exist_ok=True)

    TP = blueprint.get("TITLE_PAGE") or {}
    sys_name = TP.get("SYSTEM_NAME") or "CUSTOM"
    sys_full = TP.get("SYSTEM_FULL_NAME") or system_full or sys_name
    doc_date = datetime.now().strftime("%d %B %Y")

    doc = Document()

    # Page setup (A4)
    for section in doc.sections:
        section.page_width    = Pt(595)
        section.page_height   = Pt(842)
        section.left_margin   = Inches(1.0)
        section.right_margin  = Inches(0.9)
        section.top_margin    = Inches(1.0)
        section.bottom_margin = Inches(1.0)

    _add_header_footer(doc, sys_name, sys_full)

    # Sections (each conditional on real data)
    _build_title_page(doc, TP, doc_date)

    ES = blueprint.get("EXECUTIVE_SUMMARY")
    _build_executive_summary(doc, ES or {})

    MS = blueprint.get("MIGRATION_STRATEGY")
    _build_strategy(doc, MS or {})

    caps = blueprint.get("CAPABILITY_MIGRATIONS") or []
    _build_capabilities(doc, caps)

    DM = blueprint.get("DATA_MIGRATION")
    _build_data_migration(doc, DM or {})

    TP_TEST = blueprint.get("TESTING_PLAN")
    _build_testing(doc, TP_TEST or {})

    items = blueprint.get("GO_LIVE_CHECKLIST") or []
    _build_golive(doc, items)

    refs = blueprint.get("REFERENCES") or []
    _build_references(doc, refs)

    _build_end_page(doc, sys_name, doc_date)

    filename = f"{uuid.uuid4().hex[:12]}_blueprint.docx"
    doc.save(str(DOWNLOADS_DIR / filename))
    return f"/static/downloads/{filename}"
